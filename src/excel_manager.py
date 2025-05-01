"""
Excel manager module for the WhatsApp Money Tracker Bot.
This module handles all Excel file operations for storing transaction data.
"""

import os
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
import sys
import os

# Add the parent directory to sys.path to allow imports from the src package
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Constants
DATA_DIR = os.path.join(os.path.dirname(os.path.dirname(__file__)), 'data')
TRANSACTIONS_FILE = os.path.join(DATA_DIR, 'transactions.xlsx')

# Ensure data directory exists
os.makedirs(DATA_DIR, exist_ok=True)

def initialize_excel_file():
    """
    Initialize the Excel file if it doesn't exist.
    Creates the file with appropriate headers.
    """
    if not os.path.exists(TRANSACTIONS_FILE):
        # Create a new DataFrame with the required columns
        df = pd.DataFrame(columns=[
            'Date', 'Time', 'Contact ID', 'Contact Name', 
            'Amount', 'Type', 'Notes'
        ])
        
        # Save to Excel
        df.to_excel(TRANSACTIONS_FILE, index=False)
        print(f"Created new transactions file at {TRANSACTIONS_FILE}")
    else:
        print(f"Transactions file already exists at {TRANSACTIONS_FILE}")

def add_transaction(contact_id, contact_name, amount, transaction_type='lend', notes=''):
    """
    Add a new transaction to the Excel file.
    
    Args:
        contact_id (int): ID of the contact
        contact_name (str): Name of the contact
        amount (float): Amount of money involved
        transaction_type (str): 'lend' or 'borrow'
        notes (str): Any additional notes
        
    Returns:
        bool: True if successful, False otherwise
    """
    try:
        # Ensure file exists
        initialize_excel_file()
        
        # Current date and time
        now = datetime.now()
        date_str = now.strftime('%Y-%m-%d')
        time_str = now.strftime('%H:%M:%S')
        
        # Load existing data
        df = pd.read_excel(TRANSACTIONS_FILE)
        
        # Create new row
        new_row = {
            'Date': date_str,
            'Time': time_str,
            'Contact ID': contact_id,
            'Contact Name': contact_name,
            'Amount': float(amount),
            'Type': transaction_type,
            'Notes': notes
        }
        
        # Append new row
        df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
        
        # Save back to Excel
        df.to_excel(TRANSACTIONS_FILE, index=False)
        
        return True
    except Exception as e:
        print(f"Error adding transaction: {e}")
        return False

def generate_daily_report(date=None):
    """
    Generate a daily report for a specific date.
    
    Args:
        date (str, optional): Date in 'YYYY-MM-DD' format. Defaults to today.
        
    Returns:
        str: Path to the generated report file
    """
    try:
        # Use today's date if not specified
        if date is None:
            date = datetime.now().strftime('%Y-%m-%d')
            
        # Ensure file exists
        initialize_excel_file()
        
        # Load data
        df = pd.read_excel(TRANSACTIONS_FILE)
        
        # Filter for the specified date
        daily_df = df[df['Date'] == date]
        
        if daily_df.empty:
            print(f"No transactions found for {date}")
            return None
            
        # Create a new Excel file for the report
        report_file = os.path.join(DATA_DIR, f'daily_report_{date}.xlsx')
        
        # Create a summary
        summary = {
            'Total Lent': daily_df[daily_df['Type'] == 'lend']['Amount'].sum(),
            'Total Borrowed': daily_df[daily_df['Type'] == 'borrow']['Amount'].sum(),
            'Net Amount': daily_df[daily_df['Type'] == 'lend']['Amount'].sum() - 
                         daily_df[daily_df['Type'] == 'borrow']['Amount'].sum()
        }
        
        # Create a writer object
        with pd.ExcelWriter(report_file, engine='openpyxl') as writer:
            # Write the transactions
            daily_df.to_excel(writer, sheet_name='Transactions', index=False)
            
            # Write the summary
            summary_df = pd.DataFrame(list(summary.items()), columns=['Metric', 'Amount'])
            summary_df.to_excel(writer, sheet_name='Summary', index=False)
            
            # Access the workbook and format it
            workbook = writer.book
            
            # Format the summary sheet
            summary_sheet = workbook['Summary']
            for row in summary_sheet.iter_rows(min_row=2, max_row=4, min_col=2, max_col=2):
                for cell in row:
                    cell.number_format = '₹#,##0.00'
            
            # Format the transactions sheet
            trans_sheet = workbook['Transactions']
            for row in trans_sheet.iter_rows(min_row=2, min_col=5, max_col=5):
                for cell in row:
                    cell.number_format = '₹#,##0.00'
        
        print(f"Generated daily report at {report_file}")
        return report_file
    except Exception as e:
        print(f"Error generating daily report: {e}")
        return None

def get_balance_for_contact(contact_id):
    """
    Calculate the current balance for a specific contact.
    
    Args:
        contact_id (int): ID of the contact
        
    Returns:
        float: Current balance (positive means they owe you, negative means you owe them)
    """
    try:
        # Ensure file exists
        initialize_excel_file()
        
        # Load data
        df = pd.read_excel(TRANSACTIONS_FILE)
        
        # Filter for the specified contact
        contact_df = df[df['Contact ID'] == contact_id]
        
        if contact_df.empty:
            return 0.0
            
        # Calculate balance
        lent = contact_df[contact_df['Type'] == 'lend']['Amount'].sum()
        borrowed = contact_df[contact_df['Type'] == 'borrow']['Amount'].sum()
        
        return lent - borrowed
    except Exception as e:
        print(f"Error calculating balance: {e}")
        return 0.0

# Initialize the Excel file when the module is imported
initialize_excel_file()
